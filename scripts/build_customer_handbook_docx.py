#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Сборка руководства для заказчика (DOCX) со скриншотами из docs/customer-screenshots-2026-04-21/.
Формат страницы: ISO A4 (210×297 мм), книжная ориентация, поля 2 см.
Зависимости:
  pip install python-docx
  (рекомендуется) pip install Pillow — скриншоты сжимаются перед вставкой, файл DOCX заметно легче.

Запуск из корня репозитория:
  python scripts/build_customer_handbook_docx.py
  python scripts/build_customer_handbook_docx.py --out docs/Rukovodstvo-dlya-zakazchika.docx
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Inches, Pt
except ImportError:
    print("Установите: pip install python-docx", file=sys.stderr)
    sys.exit(1)

try:
    from io import BytesIO

    from PIL import Image
except ImportError:
    Image = None  # type: ignore[misc, assignment]
    BytesIO = None  # type: ignore[misc, assignment]


def _repo_root() -> Path:
    return Path(__file__).resolve().parent.parent


def _screenshot_dir(root: Path) -> Path:
    return root / "docs" / "customer-screenshots-2026-04-21"


def _shot_path(shots: Path, filename: str) -> Path | None:
    direct = shots / filename
    if direct.is_file():
        return direct
    if filename == "home-desktop.png":
        candidates = sorted(shots.glob("home-desktop*.png"))
        if candidates:
            return candidates[0]
    return None


def _apply_a4_portrait(section) -> None:
    """Формат страницы ISO A4 (210×297 мм), книжная ориентация, поля по 2 см."""
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)


def _add_para(doc: Document, text: str, *, bold: bool = False) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(11)


def _add_image(doc: Document, path: Path, caption: str, width_inches: float = 6.2) -> None:
    if not path.is_file():
        _add_para(doc, f"[Скриншот не найден: {path.name}]", bold=True)
        return
    img_path = str(path)
    if Image is not None and BytesIO is not None:
        try:
            im = Image.open(path).convert("RGB")
            w, h = im.size
            max_w = 1200
            if w > max_w:
                new_h = int(h * (max_w / w))
                try:
                    resample = Image.Resampling.LANCZOS
                except AttributeError:
                    resample = Image.LANCZOS  # type: ignore[attr-defined]
                im = im.resize((max_w, new_h), resample)
            buf = BytesIO()
            im.save(buf, format="JPEG", quality=82, optimize=True)
            buf.seek(0)
            doc.add_picture(buf, width=Inches(width_inches))
        except OSError:
            doc.add_picture(img_path, width=Inches(width_inches))
    else:
        doc.add_picture(img_path, width=Inches(width_inches))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True
        run.font.size = Pt(9)
    doc.add_paragraph()


def build_doc(out_path: Path) -> None:
    root = _repo_root()
    shots = _screenshot_dir(root)

    doc = Document()
    sect = doc.sections[0]
    _apply_a4_portrait(sect)

    title = doc.add_heading("Руководство по сайту «Фабрика тентов»", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(
        "Для заказчика: панель /staff, разделы сайта и скриншоты витрины\n(без технических терминов веб-разработки)"
    )
    r.italic = True
    r.font.size = Pt(12)
    doc.add_paragraph()

    doc.add_heading("1. Для кого эта памятка", level=1)
    _add_para(
        doc,
        "Документ написан простым языком. Вам не нужно разбираться в программировании, "
        "хостинге и базах данных — достаточно понимать, как пользоваться рабочей панелью и что видит клиент на сайте.",
    )

    doc.add_heading("2. Рабочая панель для сотрудников (/staff)", level=1)
    _add_para(
        doc,
        "Для обновления каталога, заявок, отзывов и материалов сайта используется отдельная панель в браузере. "
        "Пример адреса: https://fabrika-tentov.ru/staff (у вас может быть тот же путь на вашем домене).",
    )
    _add_para(
        doc,
        "Вход: логин и пароль выдаёт исполнитель проекта. После входа слева открывается меню разделов; "
        "вверху можно выйти из системы или вернуться на главный экран панели.",
    )
    _add_para(
        doc,
        "Важно: эта панель предназначена для каталога, заказов, контента и заявок. "
        "Сложные системные настройки, интеграции и учётные записи администраторов находятся в другой "
        "«технической» админке (обычно адрес вида /admin/) — её использует технический специалист.",
    )

    doc.add_heading("3. Товары и категории каталога", level=1)
    _add_para(doc, "Категории", bold=True)
    _add_para(
        doc,
        "В меню слева откройте раздел «Категории каталога». Здесь список разделов витрины. "
        "Кнопка создания добавляет новую категорию; у существующей строки нажмите значок карандаша, чтобы изменить название, "
        "описание, порядок показа или картинку.",
    )
    _add_para(doc, "Новый товар", bold=True)
    _add_para(
        doc,
        "Раздел «Товары» → кнопка создания новой записи. Заполните название, привязку к категории, описание, "
        "цену и прочие поля по подсказкам на экране. Сохраните запись — товар появится на сайте, если для него включён показ на витрине.",
    )
    _add_para(doc, "Изменить существующий товар", bold=True)
    _add_para(
        doc,
        "Список «Товары» → найдите строку с нужным названием → откройте её (карандаш или название). "
        "После правок нажмите «Сохранить». Фотографии товара, варианты (размеры, комплектации) и подробные характеристики "
        "часто задаются отдельными вкладками или связанными блоками внутри карточки товара — ориентируйтесь на заголовки полей.",
    )

    doc.add_heading("4. Акции на сайте (раздел «Акции» для покупателей)", level=1)
    _add_para(
        doc,
        "Страницы акций и их условия на витрине настраиваются в основной Django-админке, а не в панели /staff. "
        "Обычно это адрес вида https://fabrika-tentov.ru/admin/ — отдельный вход и права доступа. "
        "Если вам нужно самостоятельно вести акции, попросите исполнителя выдать учётную запись и короткую инструкцию по разделу «Акции».",
    )

    doc.add_heading("5. Отзывы", level=1)
    _add_para(
        doc,
        "Раздел «Отзывы» — полный список отзывов: можно добавить отзыв вручную или открыть существующий и изменить текст, рейтинг, "
        "фото, признак публикации на сайте.",
    )
    _add_para(
        doc,
        "«Очередь модерации отзывов» — сюда попадают отзывы, оставленные посетителями с сайта до публикации. "
        "Просмотрите текст, примите решение: опубликовать или отклонить. Так вы контролируете, что уходит на витрину.",
    )

    doc.add_heading("6. Другие разделы панели /staff", level=1)
    staff_blocks = [
        "«Главная страница» — блоки и тексты главной экрана сайта по разделам (откройте нужный блок и редактируйте поля).",
        "«Блог» — статьи: создание, правка, снятие с публикации.",
        "«Портфолио» — примеры работ с фото и описаниями.",
        "«Заказы» — заказы покупателей: просмотр состава, статусов, комментариев для менеджера.",
        "«Заявки: обратный звонок» и «Заявки: калькулятор» — обращения с форм сайта, их обработка.",
        "«Профили покупателей» и «Адреса доставки» — справочная информация по клиентам при необходимости.",
    ]
    for line in staff_blocks:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("7. Основные разделы сайта для посетителя", level=1)
    _add_para(doc, "В меню самого сайта для клиентов обычно доступно:", bold=True)
    items = [
        "Главная — первый экран: ключевое предложение, кнопки «связаться» и блоки с преимуществами.",
        "Каталог — товары с фото и ценами, можно открыть карточку и узнать подробности.",
        "Портфолио — выполненные работы, часто с фото «до и после».",
        "Блог — статьи и новости компании.",
        "Акции — специальные предложения.",
        "Контакты — телефон, адрес, карта, форма обратной связи.",
        "Отзывы — мнения клиентов.",
    ]
    for line in items:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("8. Дизайн и «настроение» сайта", level=1)
    _add_para(
        doc,
        "Внешний вид подобран так, чтобы сайт выглядел современно и надёжно: спокойные цвета, "
        "крупные заголовки, понятные кнопки. На телефоне блоки складываются в колонку — "
        "ничего не нужно уменьшать пальцами, текст читается без лупы.",
    )
    _add_para(
        doc,
        "Есть плавные подсветки при наведении и аккуратные анимации. Если у посетителя в системе "
        "включено «уменьшить движение», лишние анимации отключаются — это норма доступности.",
    )

    doc.add_heading("9. Полезные возможности для посетителя сайта", level=1)
    feats = [
        "На главной можно быстро оставить заявку (заказ обратного звонка или сообщение) — не обязательно сразу оформлять покупку.",
        "Корзина и оформление заказа ведут по шагам: что выбрали, куда доставить, как оплатить.",
        "Есть поиск по каталогу и фильтры — чтобы быстрее найти нужный тип тента или маркизы.",
        "Личный кабинет — для зарегистрированных клиентов: история заказов, адреса, смена пароля. Регистрация по желанию.",
        "Согласие с cookies и политика персональных данных — чтобы посетитель понимал, как обрабатываются данные; баннер можно принять и продолжить просмотр.",
    ]
    for line in feats:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("10. Заказ, доставка и оплата (без технических подробностей)", level=1)
    _add_para(
        doc,
        "Посетитель может положить товары в корзину и перейти к оформлению. Дальше сайт задаёт понятные вопросы: "
        "контакты, способ получения (например, курьер или пункт выдачи), способ оплаты. Конкретные варианты доставки и оплаты "
        "зависят от настроек магазина — их можно уточнить у исполнителя проекта.",
    )
    _add_para(
        doc,
        "После оформления клиент обычно получает подтверждение на электронную почту (если адрес указан). Статус заказа "
        "можно смотреть в личном кабинете.",
    )

    doc.add_heading("11. Скриншоты (как выглядит сайт для клиента)", level=1)
    _add_para(
        doc,
        "Ниже — примеры экранов с компьютера и с телефона. У вас на сайте тексты и фото могут отличаться — "
        "это нормально: контент настраивается под вашу компанию.",
    )

    pairs = [
        ("home-desktop.png", "Главная страница — вид с компьютера"),
        ("home-mobile.png", "Главная страница — вид с телефона"),
        ("catalog-desktop.png", "Каталог — компьютер"),
        ("catalog-mobile.png", "Каталог — телефон"),
        ("portfolio-desktop.png", "Портфолио — компьютер"),
        ("portfolio-mobile.png", "Портфолио — телефон"),
        ("contacts-desktop.png", "Контакты — компьютер"),
        ("contacts-mobile.png", "Контакты — телефон"),
    ]
    for name, cap in pairs:
        p = _shot_path(shots, name)
        if p is None:
            _add_para(doc, f"[Скриншот не найден: {name}]", bold=True)
            continue
        _add_image(doc, p, cap)

    doc.add_heading("12. Если что-то не работает", level=1)
    _add_para(
        doc,
        "Сначала обновите страницу (кнопка обновления в браузере). Попробуйте другой браузер или сеть Wi‑Fi/мобильный интернет. "
        "Если проблема не исчезла — зафиксируйте, что именно вы делали и что на экране, и напишите контактному лицу от исполнителя.",
    )

    doc.add_heading("13. Краткий чек-лист", level=1)
    chk = [
        "Сайт открывается по известному адресу с телефона и с компьютера.",
        "В каталоге видны актуальные товары и цены (если цены включены).",
        "Контакты и реквизиты совпадают с вашими действующими.",
        "Форма обратной связи и/или заказа доходит до ответственного сотрудника.",
        "Панель /staff открывается под выданной учётной записью, товары и отзывы сохраняются без ошибок.",
    ]
    for line in chk:
        doc.add_paragraph(line, style="List Number")

    doc.add_paragraph()
    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = foot.add_run(
        "Документ сформирован автоматически из материалов проекта.\n"
        "При обновлении дизайна скриншоты можно заменить в папке docs/customer-screenshots-2026-04-21 "
        "и заново запустить: python scripts/build_customer_handbook_docx.py"
    )
    rr.italic = True
    rr.font.size = Pt(9)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"Готово: {out_path}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--out",
        type=Path,
        default=_repo_root() / "docs" / "Rukovodstvo-dlya-zakazchika.docx",
        help="Путь к выходному DOCX",
    )
    args = ap.parse_args()
    build_doc(args.out.resolve())


if __name__ == "__main__":
    main()
